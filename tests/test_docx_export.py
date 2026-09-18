import os

from docx import Document

from aquaculture_manuscript_mcp.tools import docx_export


def test_export_creates_real_readable_docx(tmp_path):
    output_path = str(tmp_path / "manuscript.docx")

    result = docx_export.export_manuscript_docx(
        output_path=output_path,
        title="Effects of Dietary Protein and Feeding Frequency on Growth in Nile Tilapia",
        abstract="A study was conducted to evaluate growth and water quality outcomes.",
        keywords=["Oreochromis niloticus", "feeding frequency", "water quality"],
        sections=[
            {"heading": "Introduction", "body": "The problem statement.\n\nA second paragraph."},
            {"heading": "Materials and Methods", "body": "Fish, system, design."},
            {"heading": "Results", "body": "SGR differed significantly (P<0.001)."},
        ],
        authors=["A. Researcher", "B. Scientist"],
        tables=[
            {
                "caption": "Table 1. Growth performance by treatment.",
                "headers": ["Treatment", "SGR", "FCR"],
                "rows": [["T1", "2.05", "1.68"], ["T2", "2.48", "1.30"]],
            }
        ],
        references=["Yossa, R. (2014). Writing a Scientific Manuscript. J. Appl. Aquacult. 26:293-309."],
        ai_disclosure="During preparation, the author(s) used [TOOL] to assist with X.",
    )

    assert os.path.exists(output_path)
    assert result["file_path"] == output_path
    assert result["table_count"] == 1
    assert result["reference_count"] == 1
    assert result["has_ai_disclosure"] is True

    # Read it back with python-docx to confirm the content actually landed,
    # not just that a file exists.
    doc = Document(output_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Effects of Dietary Protein" in full_text
    assert "A. Researcher, B. Scientist" in full_text
    assert "A study was conducted" in full_text
    assert "Oreochromis niloticus" in full_text
    assert "The problem statement." in full_text
    assert "SGR differed significantly" in full_text
    assert "Declaration of generative AI" in full_text
    assert "During preparation" in full_text
    assert "Yossa, R. (2014)" in full_text

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Treatment", "SGR", "FCR"]
    assert [c.text for c in table.rows[1].cells] == ["T1", "2.05", "1.68"]


def test_export_without_optional_fields_still_works(tmp_path):
    output_path = str(tmp_path / "minimal.docx")
    result = docx_export.export_manuscript_docx(
        output_path=output_path,
        title="Minimal Manuscript",
        abstract="Short abstract.",
        keywords=[],
        sections=[{"heading": "Introduction", "body": "Text."}],
    )
    assert os.path.exists(output_path)
    assert result["table_count"] == 0
    assert result["reference_count"] == 0
    assert result["has_ai_disclosure"] is False

    doc = Document(output_path)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Declaration of generative AI" not in full_text


def test_word_counts_reflect_actual_content(tmp_path):
    output_path = str(tmp_path / "counts.docx")
    result = docx_export.export_manuscript_docx(
        output_path=output_path,
        title="T",
        abstract="one two three four five",
        keywords=[],
        sections=[{"heading": "Introduction", "body": "one two three"}],
    )
    assert result["word_counts"]["Abstract"] == 5
    assert result["word_counts"]["Introduction"] == 3
