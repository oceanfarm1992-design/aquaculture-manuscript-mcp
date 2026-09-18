"""Assembles a real .docx manuscript file from structured content.

This is a rendering tool, not a writing tool: it lays out whatever title,
abstract, sections, tables, and references it's given into a real Word
document with standard academic manuscript formatting. It never invents
content — a missing section, empty references list, or placeholder text is
passed through as given, not filled in with something plausible-sounding.
"""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _set_academic_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def _add_table(document: Document, caption: str, headers: list[str], rows: list[list]) -> None:
    caption_p = document.add_paragraph()
    caption_run = caption_p.add_run(caption)
    caption_run.bold = True
    caption_run.font.size = Pt(11)

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = str(header)
        for p in header_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if value is None else str(value)

    document.add_paragraph()


def export_manuscript_docx(
    output_path: str,
    title: str,
    abstract: str,
    keywords: list[str],
    sections: list[dict],
    authors: list[str] | None = None,
    tables: list[dict] | None = None,
    references: list[str] | None = None,
    ai_disclosure: str | None = None,
) -> dict:
    """Write a real .docx manuscript file to output_path.

    sections: ordered list of {"heading": str, "body": str}, e.g.
        [{"heading": "Introduction", "body": "..."},
         {"heading": "Materials and Methods", "body": "..."}, ...]
    tables: optional list of {"caption": str, "headers": [str, ...],
        "rows": [[...], ...]}, inserted in order after the sections.
    references: optional list of pre-formatted citation strings (this tool
        does not format or verify citations — see resolve_doi_metadata /
        validate_bibtex for that).
    ai_disclosure: optional exact declaration text (see draft_ai_disclosure)
        to place in its own section before the references list.

    Returns file path, word counts per section, and table/reference counts —
    not a judgment on whether the content is publication-ready.
    """
    document = Document()
    _set_academic_style(document)

    title_p = document.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(14)

    if authors:
        authors_p = document.add_paragraph()
        authors_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        authors_p.add_run(", ".join(authors)).italic = True

    document.add_paragraph()

    abstract_heading = document.add_paragraph()
    abstract_heading.add_run("Abstract").bold = True
    document.add_paragraph(abstract)

    if keywords:
        keywords_p = document.add_paragraph()
        keywords_p.add_run("Keywords: ").bold = True
        keywords_p.add_run(", ".join(keywords))

    document.add_page_break()

    word_counts = {"Abstract": len(abstract.split())}

    for section in sections:
        heading = section.get("heading", "")
        body = section.get("body", "")
        document.add_heading(heading, level=1)
        for paragraph_text in body.split("\n\n"):
            if paragraph_text.strip():
                document.add_paragraph(paragraph_text.strip())
        word_counts[heading] = len(body.split())

    for table_spec in tables or []:
        _add_table(
            document,
            table_spec.get("caption", ""),
            table_spec.get("headers", []),
            table_spec.get("rows", []),
        )

    if ai_disclosure:
        document.add_heading(
            "Declaration of generative AI and AI-assisted technologies in the manuscript preparation process",
            level=1,
        )
        document.add_paragraph(ai_disclosure)

    if references:
        document.add_heading("References", level=1)
        for ref in references:
            p = document.add_paragraph(ref)
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)

    document.save(output_path)

    return {
        "file_path": output_path,
        "word_counts": word_counts,
        "table_count": len(tables or []),
        "reference_count": len(references or []),
        "has_ai_disclosure": bool(ai_disclosure),
    }
